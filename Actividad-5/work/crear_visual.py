from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image
import pypdfium2 as pdfium
import re, json

ROOT=Path(r'C:\Users\Javier\Documents\GitHub\Actividad-5')
OUT=ROOT/'outputs'; WORK=ROOT/'work'
PPT=OUT/'Presentacion_Moda_Urbana_visual.pptx'; PDF=OUT/'Presentacion_Moda_Urbana_visual.pdf'
BG='F5F1E9'; INK='192E30'; CORAL='D65A3A'; GREEN='34675C'; MUTED='64716C'; WHITE='FFFFFF'; PALE='DFE8DD'; SAND='E8DAC7'; LINE='D6D9CD'
pdfmetrics.registerFont(TTFont('Arial',r'C:\Windows\Fonts\arial.ttf'))
pdfmetrics.registerFont(TTFont('ArialBold',r'C:\Windows\Fonts\arialbd.ttf'))
prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
cv=canvas.Canvas(str(PDF),pagesize=(960,540)); cv.setTitle('Moda Urbana · Antes y después · Presentación visual')
slide=None
def rgb(h):return tuple(int(h[i:i+2],16)/255 for i in (0,2,4))
def col(h):return RGBColor.from_string(h)
def shape(x,y,w,h,fill,kind=MSO_SHAPE.RECTANGLE,r=0):
 s=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));s.fill.solid();s.fill.fore_color.rgb=col(fill);s.line.fill.background()
 if kind==MSO_SHAPE.ROUNDED_RECTANGLE:s.adjustments[0]=.13
 cv.setFillColorRGB(*rgb(fill))
 if kind==MSO_SHAPE.OVAL:cv.ellipse(x*72,540-(y+h)*72,(x+w)*72,540-y*72,stroke=0,fill=1)
 else:cv.roundRect(x*72,540-(y+h)*72,w*72,h*72,r*72,stroke=0,fill=1)
 return s
def box(x,y,w,h,fill=WHITE,r=0):return shape(x,y,w,h,fill,MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,r)
def dot(x,y,d,fill):return shape(x,y,d,d,fill,MSO_SHAPE.OVAL)
def line(x,y,x2,y2,fill=INK,width=2):
 s=slide.shapes.add_connector(1,Inches(x),Inches(y),Inches(x2),Inches(y2));s.line.color.rgb=col(fill);s.line.width=Pt(width)
 cv.setStrokeColorRGB(*rgb(fill));cv.setLineWidth(width);cv.line(x*72,540-y*72,x2*72,540-y2*72)
def txt(x,y,w,h,t,size=26,fill=INK,bold=False,center=False):
 s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=s.text_frame;tf.word_wrap=False
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,l in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=l;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=col(fill);p.space_after=Pt(0);p.line_spacing=1.08
  if center:p.alignment=PP_ALIGN.CENTER
 cv.setFillColorRGB(*rgb(fill));cv.setFont('ArialBold' if bold else 'Arial',size)
 for i,l in enumerate(t.split('\n')):
  yy=540-y*72-size*.94-i*size*1.08
  if center:cv.drawCentredString((x+w/2)*72,yy,l)
  else:cv.drawString(x*72,yy,l)
 return s
def poly(points,fill):
 builder=slide.shapes.build_freeform(points[0][0]*1000,points[0][1]*1000,scale=Inches(.001))
 builder.add_line_segments([(x*1000,y*1000) for x,y in points[1:]],close=True)
 s=builder.convert_to_shape();s.fill.solid();s.fill.fore_color.rgb=col(fill);s.line.fill.background()
 p=cv.beginPath();p.moveTo(points[0][0]*72,540-points[0][1]*72)
 for x,y in points[1:]:p.lineTo(x*72,540-y*72)
 p.close();cv.setFillColorRGB(*rgb(fill));cv.drawPath(p,stroke=0,fill=1)
 return s
def arrow(x,y,length=1.0,fill=CORAL):
 line(x,y,x+length-.1,y,fill,3);poly([(x+length-.3,y-.13),(x+length,y),(x+length-.3,y+.13)],fill)
def label(x,y,w,t,fill=CORAL):box(x,y,w,.37,fill,r=.08);txt(x+.08,y+.055,w-.16,.25,t,12,WHITE,True,True)
def new(n,title='',dark=False):
 global slide
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=col(INK if dark else BG)
 cv.setFillColorRGB(*rgb(INK if dark else BG));cv.rect(0,0,960,540,stroke=0,fill=1)
 txt(.6,.32,5,.23,'MODA URBANA',10,PALE if dark else MUTED,True)
 txt(12.15,.32,.55,.23,f'{n:02}',10,PALE if dark else MUTED,True,True)
 if title:txt(.6,.94,12.1,.75,title,34,WHITE if dark else INK,True)
 txt(.6,7.07,4,.23,'Caso ficticio',9,PALE if dark else MUTED)
def end(notes):slide.notes_slide.notes_text_frame.text=notes;cv.showPage()
def tee(x,y,w,fill=GREEN):
 pts=[(.09,.20),(.34,.04),(.43,.12),(.57,.12),(.66,.04),(.91,.20),(1,.41),(.79,.50),(.73,.37),(.73,1),(.27,1),(.27,.37),(.21,.50),(0,.41)]
 poly([(x+a*w,y+b*w) for a,b in pts],fill)
 line(x+.43*w,y+.13*w,x+.57*w,y+.13*w,BG,2)
def store(x,y,w=3.4):
 h=w*.74;box(x,y,w,h,PALE,r=.16)
 box(x+.16*w,y+.14*w,.68*w,.16*w,GREEN)
 for i in range(6):box(x+(.16+i*.113)*w,y+.14*w,.056*w,.16*w,CORAL)
 box(x+.13*w,y+.36*w,.42*w,.29*w,WHITE);box(x+.64*w,y+.35*w,.22*w,.39*w,INK)
 tee(x+.20*w,y+.40*w,.24*w,GREEN)
def notebook(x,y,w=2.8,h=2.9):
 box(x,y,w,h,SAND,r=.13);box(x+.14,y+.12,w-.25,h-.22,WHITE,r=.06)
 for j in range(4):dot(x-.04,y+.4+j*.56,.12,INK)
 for j in range(4):line(x+.55,y+.55+j*.5,x+w-.35,y+.55+j*.5,LINE,2)
 for j in range(3):line(x+.7+j*.22,y+1.03,x+.7+j*.22,y+1.42,CORAL,2)
 line(x+.64,y+1.39,x+1.36,y+1.04,CORAL,2)
def clock(x,y,d=1.1):
 dot(x,y,d,INK);dot(x+.07,y+.07,d-.14,BG);line(x+d/2,y+d/2,x+d/2,y+.2*d,INK,3);line(x+d/2,y+d/2,x+.76*d,y+.65*d,INK,3)
def browser(x,y,w=5.8,h=3.7):
 box(x,y,w,h,WHITE,r=.12);box(x,y,w,h*.11,INK)
 for j in range(3):dot(x+w*(.03+j*.028),y+h*.035,w*.014,PALE)
 box(x+w*.045,y+h*.18,w*.37,h*.72,PALE,r=.04);tee(x+w*.075,y+h*.28,w*.30,GREEN)
 for j,c in enumerate([GREEN,CORAL,MUTED]):
  xx=x+w*.48+j*w*.16;box(xx,y+h*.19,w*.13,h*.28,'EAE6DB',r=.04);tee(xx+w*.01,y+h*.22,w*.105,c)
 box(x+w*.48,y+h*.54,w*.4,h*.027,LINE);box(x+w*.48,y+h*.61,w*.28,h*.027,LINE)
 for j,s in enumerate(['S','M','L']):
  xx=x+w*(.48+j*.085);box(xx,y+h*.71,w*.065,h*.08,PALE);txt(xx,y+h*.725,w*.065,h*.07,s,7 if w<4 else 10,INK,True,True)
 box(x+w*.48,y+h*.84,w*.40,h*.08,CORAL,r=.05)
def check(x,y,size=.6,fill=GREEN):line(x,y+size*.45,x+size*.32,y+size*.78,fill,5);line(x+size*.32,y+size*.78,x+size,y,fill,5)
def carton(x,y,w=2.4):
 box(x,y,w,w*.7,SAND);poly([(x,y),(x+w*.22,y-w*.22),(x+w*1.06,y-w*.22),(x+w,y)],'D0BE9F')
 box(x+w*.43,y,w*.14,w*.7,WHITE);box(x+w*.12,y+w*.28,w*.24,w*.18,WHITE)
def calendar(x,y,w=2.5):
 box(x,y,w,w*.85,WHITE,r=.1);box(x,y,w,w*.20,GREEN)
 for a in range(4):
  for b in range(3):box(x+w*(.10+a*.22),y+w*(.32+b*.17),w*.09,w*.08,PALE)
 box(x+w*.32,y+w*.49,w*.09,w*.08,CORAL)
def person(x,y,s=1,fill=GREEN):dot(x+.25*s,y,.5*s,fill);box(x,y+.60*s,s,.95*s,fill,r=.2)

# 01: one image transforms into the other.
new(1,dark=True)
txt(.65,1.02,10.8,1.3,'De la libreta\na la tienda digital.',43,WHITE,True)
notebook(1.45,3.05,2.25,2.55);arrow(4.1,4.2,1.7,CORAL);browser(6.4,2.9,5.8,3.7)
end('Presentamos la transformación digital de Moda Urbana, una tienda de ropa ficticia. Partimos de un comercio que organiza sus tareas de forma manual. La idea no es solo tener una página web: queremos mostrar cómo conectar las ventas, las existencias y los pedidos puede facilitar el trabajo y la atención al cliente. Los resultados que comentamos son esperados, no mediciones reales.')

new(2,'El recorrido')
store(.9,2.8,3.0);arrow(4.0,3.85,.85);browser(5.0,2.65,3.1,2.7);arrow(8.35,3.85,.85);dot(9.65,2.9,2.6,PALE);check(10.34,3.75,1.15)
txt(.9,5.85,3,.5,'ANTES',24,INK,True,True);txt(5.0,5.85,3.1,.5,'CAMBIO',24,CORAL,True,True);txt(9.4,5.85,3.1,.5,'MEJORA',24,GREEN,True,True)
end('El recorrido tiene tres partes. Primero veremos cómo trabaja la tienda antes de digitalizarse. Después enseñaremos los cambios que proponemos y la web de demostración. Por último explicaremos qué problemas podrían mejorar y cómo comprobaríamos si el cambio funciona.')

new(3,'Antes: todo depende del local')
store(.85,2.6,4.25);clock(4.15,2.15,1.35);notebook(6.05,2.55,2.6,3.15);carton(9.75,3.2,2.0)
txt(.75,6.03,4.5,.45,'HORARIO LIMITADO',18,CORAL,True,True);txt(5.6,6.03,3.5,.45,'STOCK A MANO',18,CORAL,True,True);txt(9.35,6.03,3.1,.45,'ENCARGOS EN PAPEL',18,CORAL,True,True)
end('Antes de la transformación, la tienda depende de las visitas al local y de su horario. Para consultar una talla, el personal revisa las prendas y el almacén. Las existencias y los encargos se apuntan en una libreta. Esto no significa que se trabaje mal, sino que la información está dispersa y puede costar más tiempo encontrarla o actualizarla.')

new(4,'Después: la información se conecta')
store(.8,2.9,3.0);browser(5.1,2.55,3.25,2.8)
box(10.0,2.8,2.4,2.8,PALE,r=.12)
for j in range(3):box(10.28,3.1+j*.65,1.83,.42,GREEN)
arrow(3.95,3.9,.9);arrow(8.62,3.9,1.1)
txt(.8,5.95,3,.45,'TIENDA',22,INK,True,True);txt(5.15,5.95,3.2,.45,'WEB',22,INK,True,True);txt(9.75,5.95,2.9,.45,'INVENTARIO',22,INK,True,True)
end('Después, la tienda, la web y el inventario deberían compartir la misma información. Una venta tiene que quedar registrada y actualizar las existencias de la talla y el color elegidos. Así podemos evitar copiar la misma información en varios sitios y reducir el riesgo de ofrecer una prenda que ya está agotada. En una tienda real haría falta configurar y comprobar esta integración.')

new(5,'La colección, desde casa')
browser(.8,2.07,7.7,4.48)
txt(9.15,2.55,3.3,.6,'VER',30,GREEN,True);txt(9.15,3.65,3.3,.6,'ELEGIR',30,GREEN,True);txt(9.15,4.75,3.3,.6,'PEDIR',30,CORAL,True)
end('La página web es la parte más visible de la propuesta. El cliente puede consultar la colección, los precios y la disponibilidad sin desplazarse al local. En nuestra demostración puede buscar y filtrar prendas, elegir talla y color y añadirlas a una cesta. El pedido es simulado: no se realiza ningún cobro ni se envía mercancía.')

new(6,'Saber qué queda')
label(.9,2.07,1.4,'ANTES');notebook(1.05,2.75,3.0,2.92)
arrow(4.7,4.05,1.15)
label(6.45,2.07,1.6,'DESPUÉS',GREEN)
box(6.45,2.75,5.75,2.92,WHITE,r=.12);tee(6.78,3.25,1.25,GREEN)
for j,(s,n) in enumerate([('S','3'),('M','4'),('L','0')]):
 xx=8.55+j*1.1;txt(xx,3.25,.85,.5,s,27,MUTED,True,True);txt(xx,4.1,.85,.7,n,42,CORAL if n=='0' else GREEN,True,True)
txt(6.45,6.1,5.75,.45,'TALLA + COLOR + UNIDADES',20,GREEN,True,True)
end('El cambio en el inventario consiste en pasar de una anotación manual a un registro por prenda, talla y color. La imagen muestra unidades de ejemplo. En la web hay una combinación agotada para que se pueda comprobar cómo se bloquea la compra cuando no queda stock. La mejora esperada es consultar más rápido y cometer menos errores, siempre que las entradas, ventas y devoluciones se registren correctamente.')

new(7,'Cada pedido tiene un recorrido')
browser(.85,2.8,3.1,2.75);arrow(4.1,3.95,.65);carton(5.2,3.35,2.4);arrow(8.05,3.95,.65);dot(9.2,2.8,2.7,PALE);check(9.87,3.65,1.25)
txt(.7,5.95,3.5,.45,'RECIBIDO',22,INK,True,True);txt(4.85,5.95,3.2,.45,'PREPARACIÓN',22,INK,True,True);txt(8.95,5.95,3.3,.45,'PREPARADO',22,GREEN,True,True)
end('Antes, un encargo anotado en papel puede perderse o no dejar claro quién debe prepararlo. Con un registro digital, el pedido tiene un estado. En la demostración aparece primero como pendiente y después se puede marcar como preparado. Esto permite enseñar cómo se organiza el seguimiento. La entrega y el envío real no están incluidos en esta web de prueba.')

new(8,'Un equipo más organizado')
person(1.3,2.55,1.5,GREEN);person(3.1,3.55,1.2,CORAL);arrow(4.9,3.85,1.15);calendar(7.05,2.3,4.25)
txt(1.0,6.1,4.2,.5,'EQUIPO',24,INK,True,True);txt(7.05,6.1,4.25,.5,'TURNOS + TAREAS',24,GREEN,True,True)
end('La transformación también afecta a las personas. Un calendario compartido permitiría consultar los turnos, las vacaciones y las tareas previstas. Por ejemplo, quién recibe la mercancía y quién prepara los pedidos. Para que funcione, el equipo necesitaría una formación inicial y acuerdos sobre cómo actualizar la información. Este calendario forma parte de la propuesta del informe, no de la web de demostración.')

new(9,'Una herramienta para cada tarea')
for x,c in [(1.15,GREEN),(5.15,CORAL),(9.15,INK)]:dot(x,2.55,2.4,c)
txt(1.15,3.25,2.4,.6,'O',48,WHITE,True,True);txt(5.15,3.25,2.4,.6,'C',48,WHITE,True,True)
calendar(9.65,3.15,1.4)
txt(.8,5.35,3.1,.5,'Odoo',28,INK,True,True);txt(4.8,5.35,3.1,.5,'Canva',28,INK,True,True);txt(8.8,5.35,3.1,.5,'Calendario',28,INK,True,True)
txt(.8,6.0,3.1,.4,'GESTIÓN',15,GREEN,True,True);txt(4.8,6.0,3.1,.4,'COMUNICACIÓN',15,CORAL,True,True);txt(8.8,6.0,3.1,.4,'COORDINACIÓN',15,INK,True,True)
end('Proponemos Odoo para integrar punto de venta, inventario y comercio electrónico. Su documentación oficial contempla variantes de productos y la conexión entre ventas e inventario. Canva se utilizaría para las publicaciones de novedades y para la presentación. Un calendario compartido ayudaría a organizar al personal. Los círculos de esta diapositiva son símbolos ilustrativos, no logotipos oficiales. La elección final dependería de las necesidades y el presupuesto de la tienda.')

new(10,'La prueba del cambio')
browser(.85,2.25,5.5,3.75);arrow(6.55,4.1,.9)
box(7.8,2.7,4.5,2.85,PALE,r=.12)
txt(8.05,3.1,4,.45,'PEDIDO +1',28,GREEN,True,True);line(8.22,3.95,11.87,3.95,LINE,2);txt(8.05,4.35,4,.45,'STOCK −1',28,CORAL,True,True)
txt(.85,6.35,11.5,.35,'DEMO LOCAL · SIN COBROS',16,MUTED,True,True)
end('Para enseñarlo en clase, abrimos la web y consultamos el inventario de una camiseta. Después seleccionamos esa misma talla y color, la añadimos a la cesta y confirmamos el pedido de prueba. Al volver a Gestión de prueba, aparece un nuevo pedido y una unidad menos de esa variante. Podemos marcarlo como preparado y reiniciar los datos para repetir la explicación. Todo ocurre en este navegador: no hay cobros, envíos ni conexión real con Odoo.')

new(11,'Mejoras que hay que comprobar')
clock(1.25,2.6,2.3);dot(5.5,2.6,2.3,PALE);check(6.12,3.31,1.05)
box(9.55,3.25,.45,1.65,PALE);box(10.2,2.85,.45,2.05,GREEN);box(10.85,2.25,.45,2.65,CORAL)
txt(.75,5.5,3.3,.5,'TIEMPO',24,INK,True,True);txt(4.95,5.5,3.3,.5,'ERRORES',24,INK,True,True);txt(9.0,5.5,3.3,.5,'VENTAS',24,INK,True,True)
txt(.8,6.3,11.7,.4,'COMPARAR ANTES Y DESPUÉS',18,GREEN,True,True)
end('No podemos afirmar que una tienda vende más solo por crear una web. Para valorar la mejora compararíamos el tiempo necesario para consultar una talla, las diferencias entre stock registrado y real, los errores en los pedidos y las ventas del canal online. Los periodos deberían ser parecidos y habría que tener en cuenta las rebajas o la temporada. El gráfico es una ilustración de medición, no datos reales ni una promesa de crecimiento.')

new(12,dark=True)
txt(.7,1.0,11.8,1.4,'Digitalizar es\ntrabajar mejor.',48,WHITE,True)
store(.85,3.2,2.6);arrow(3.85,4.15,.65);browser(5.0,3.1,3.2,2.65);arrow(8.65,4.15,.65);check(10.2,3.65,1.4,PALE)
txt(.7,6.25,12,.3,'FUENTES · Odoo: funciones del punto de venta · Canva: redes sociales',11,PALE)
txt(.7,6.62,12,.25,'odoo.com/es_ES/app/point-of-sale-features  ·  canva.com/es_es/redes-sociales/',9,PALE)
end('La conclusión es que la web abre una nueva forma de consultar y comprar, pero la transformación digital se consigue al conectar el trabajo de la tienda. Organizar la información puede facilitar las tareas del personal y mejorar la atención al cliente. Esta propuesta es un caso ficticio. Las fuentes de las funciones de las herramientas son las páginas oficiales de Odoo, https://www.odoo.com/es_ES/app/point-of-sale-features, y Canva, https://www.canva.com/es_es/redes-sociales/.')

cv.save();prs.save(PPT)
with ZipFile(PPT) as z:assert z.testzip() is None
checkprs=Presentation(PPT);assert len(checkprs.slides)==12
counts=[];outside=[];text_overs=[]
for i,s in enumerate(checkprs.slides,1):
 assert s.notes_slide.notes_text_frame.text.strip()
 words=[]
 for sh in s.shapes:
  if sh.left<0 or sh.top<0 or sh.left+sh.width>checkprs.slide_width+100 or sh.top+sh.height>checkprs.slide_height+100:outside.append((i,sh.name))
  if sh.has_text_frame:
   words.extend(re.findall(r'\b[\wÁÉÍÓÚáéíóúñÑ]+\b',sh.text))
   for para in sh.text_frame.paragraphs:
    if para.text:
     run=para.runs[0];size=(run.font.size or para.font.size).pt
     width=pdfmetrics.stringWidth(para.text,'ArialBold' if (run.font.bold if run.font.bold is not None else para.font.bold) else 'Arial',size)
     if width>sh.width/12700+.4:text_overs.append((i,para.text,width,sh.width/12700))
 counts.append(len(words))
assert not outside,outside
assert not text_overs,text_overs
assert max(counts[:-1])<=24,counts
doc=pdfium.PdfDocument(str(PDF));assert len(doc)==12
thumbs=[]
for i in range(len(doc)):
 im=doc[i].render(scale=.72).to_pil().convert('RGB');im.save(WORK/f'visual_slide_{i+1:02}.png');thumbs.append(im)
sw,sh=thumbs[0].size
sheet=Image.new('RGB',(sw*3,sh*4),'white')
for i,im in enumerate(thumbs):sheet.paste(im,((i%3)*sw,(i//3)*sh))
sheet.save(WORK/'visual_contacto.jpg')
(WORK/'visual_verificacion.json').write_text(json.dumps({'slides':12,'words_per_slide':counts,'outside_shapes':outside,'text_width_overflow':text_overs,'speaker_notes':12},ensure_ascii=False,indent=2),encoding='utf8')
print(PPT);print(PDF);print('Palabras por diapositiva:',counts);print('12 notas; sin elementos fuera ni líneas de texto más anchas que su caja.')
