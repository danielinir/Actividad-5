from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image, ImageOps, ImageDraw
import pypdfium2 as pdfium

ROOT=Path(r'C:\Users\Javier\Documents\GitHub\Actividad-5')
OUT=ROOT/'outputs'; WORK=ROOT/'work'; OUT.mkdir(exist_ok=True); WORK.mkdir(exist_ok=True)
PPT=OUT/'Presentacion_Moda_Urbana.pptx'; PDF=OUT/'Presentacion_Moda_Urbana.pdf'
pdfmetrics.registerFont(TTFont('Arial',r'C:\Windows\Fonts\arial.ttf'))
pdfmetrics.registerFont(TTFont('ArialBold',r'C:\Windows\Fonts\arialbd.ttf'))
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
cv=canvas.Canvas(str(PDF),pagesize=(960,540)); cv.setTitle('Moda Urbana · Transformación digital')
BG='F5F1E9'; INK='192E30'; CORAL='D65A3A'; GREEN='34675C'; MUTED='586665'; WHITE='FFFFFF'; LINE='DDD8CF'; PALE='E1EAE4'
slide=None
def color(h): return RGBColor.from_string(h)
def rgb(h): return tuple(int(h[i:i+2],16)/255 for i in (0,2,4))
def box(x,y,w,h,fill=WHITE,r=0,line=None):
 shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
 shp.fill.solid(); shp.fill.fore_color.rgb=color(fill); shp.line.fill.background()
 if r: shp.adjustments[0]=.12
 cv.setFillColorRGB(*rgb(fill)); cv.roundRect(x*72,540-(y+h)*72,w*72,h*72,r*72,stroke=0,fill=1)
 return shp
def text(x,y,w,h,txt,size=22,fill=INK,bold=False):
 tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=tb.text_frame; tf.word_wrap=True; tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
 lines=txt.split('\n')
 for i,ln in enumerate(lines):
  par=tf.paragraphs[0] if i==0 else tf.add_paragraph(); par.text=ln; par.font.name='Arial'; par.font.size=Pt(size); par.font.bold=bold; par.font.color.rgb=color(fill); par.space_after=Pt(3)
 cv.setFillColorRGB(*rgb(fill)); cv.setFont('ArialBold' if bold else 'Arial',size)
 for i,ln in enumerate(lines): cv.drawString(x*72,540-y*72-size*.91-i*(size*1.15+3),ln)
 return tb
def line(x1,y1,x2,y2,fill=INK,width=1):
 shp=slide.shapes.add_connector(1,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); shp.line.color.rgb=color(fill); shp.line.width=Pt(width)
 cv.setStrokeColorRGB(*rgb(fill)); cv.setLineWidth(width); cv.line(x1*72,540-y1*72,x2*72,540-y2*72)
def tag(x,y,txt,fill=GREEN):
 box(x,y,1.65,.35,fill,r=.1); text(x+.12,y+.065,1.45,.22,txt,10,WHITE,True)
def new(n,title,kicker='MODA URBANA / TRANSFORMACIÓN DIGITAL',dark=False):
 global slide
 slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color(INK if dark else BG)
 cv.setFillColorRGB(*rgb(INK if dark else BG)); cv.rect(0,0,960,540,stroke=0,fill=1)
 text(.65,.35,11.8,.3,kicker,10,PALE if dark else MUTED,True)
 if title: text(.65,.95,12,1.1,title,34,WHITE if dark else INK,True)
 line(.65,7.03,12.68,7.03,fill=MUTED if dark else LINE)
 text(.65,7.13,10,.2,'Caso ficticio · Propuesta académica',9,PALE if dark else MUTED)
 text(12.2,7.10,.45,.22,f'{n:02}',11,PALE if dark else MUTED,True)
def finish(notes):
 slide.notes_slide.notes_text_frame.text=notes
 cv.showPage()
def card(x,y,w,label,body,accent=GREEN):
 box(x,y,w,2.6,WHITE,r=.16); box(x,y,.07,2.6,accent)
 text(x+.25,y+.28,w-.5,.4,label,19,accent,True)
 text(x+.25,y+.95,w-.5,1.45,body,21)
def arrow(x,y): text(x,y,.55,.55,'>',32,CORAL,True)
def store(x,y,scale=1):
 box(x,y,3.2*scale,2.2*scale,PALE,r=.15)
 box(x+.25*scale,y+.38*scale,2.7*scale,.42*scale,GREEN)
 for i in range(6): box(x+(.25+i*.45)*scale,y+.38*scale,.225*scale,.42*scale,CORAL)
 box(x+.4*scale,y+.97*scale,1.35*scale,.9*scale,WHITE)
 box(x+1.95*scale,y+.97*scale,.8*scale,1.23*scale,INK)
 text(x+.45*scale,y+1.17*scale,1.3*scale,.4*scale,'MU',25*scale,GREEN,True)
def browser(x,y,w=4.8,h=3.1):
 box(x,y,w,h,WHITE,r=.12); box(x,y,w,.38,INK); text(x+.18,y+.10,w-.3,.18,'moda urbana / catálogo',9,WHITE)
 for i,c in enumerate([GREEN,CORAL,MUTED]):
  box(x+.2+i*(w-.4)/3,y+.62,(w-.65)/3,1.35,PALE if i!=1 else 'F2DED4',r=.06)
  # A simple editable garment silhouette.
  bx=x+.45+i*(w-.4)/3
  box(bx,y+.94,.55,.75,c,r=.03); box(bx-.18,y+.94,.91,.26,c,r=.03)
  text(x+.2+i*(w-.4)/3,y+2.12,(w-.5)/3,.2,['Essential','City','Straight'][i],10,INK,True)
  text(x+.2+i*(w-.4)/3,y+2.43,(w-.5)/3,.2,['19,90 €','39,90 €','49,90 €'][i],12)

new(1,'',dark=True)
tag(.7,1.15,'ANTES / DESPUÉS',CORAL)
text(.7,1.95,8.2,1.8,'Una tienda de siempre.\nUna nueva forma de vender.',38,WHITE,True)
text(.7,4.25,7.5,.8,'Moda Urbana: de la libreta al catálogo digital.',23,PALE)
store(9,3.6,.85)
finish('Vamos a explicar cómo puede mejorar una tienda de ropa al digitalizar sus procesos. Moda Urbana es un negocio ficticio. Partimos de una tienda que trabaja de forma manual y proponemos cambios concretos. No presentamos resultados reales ni damos por hecho que una web garantice vender más.')

new(2,'El recorrido del cambio')
for i,(num,title,body) in enumerate([('01','Punto de partida','Qué hace la tienda\ny qué dificultades tiene.'),('02','Transformación','Qué cambia con la web\ny la gestión digital.'),('03','Resultado esperado','Qué mejora y cómo\npodríamos comprobarlo.')]):
 x=.7+i*4.2; text(x,2.45,3,.7,num,46,CORAL,True); text(x,3.45,3.8,.55,title,25,INK,True); text(x,4.3,3.7,1.2,body,22)
finish('La presentación tiene tres partes. Primero veremos el funcionamiento inicial. Después explicaremos la web y los cambios en la gestión. Al final relacionaremos cada cambio con una mejora y diremos cómo comprobar si funciona.')

new(3,'Antes: todo depende del local')
store(.8,2.5,1.25)
text(5.6,2.35,6.8,.6,'Una pequeña tienda de ropa',26,INK,True)
for y,txt in [(3.35,'Venta durante el horario de apertura'),(4.15,'Existencias apuntadas en una libreta'),(4.95,'Encargos en papel y consultas por teléfono')]:
 box(5.65,y+.09,.12,.12,CORAL); text(5.95,y,6.4,.55,txt,21)
finish('La tienda vende ropa y complementos en su local. Para este caso suponemos que las existencias se apuntan en una libreta y los encargos en papel. Si alguien pregunta por una talla, el personal tiene que comprobarla físicamente. Esto consume tiempo y facilita que haya olvidos.')

new(4,'El problema se nota en las tareas diarias')
card(.7,2.45,3.8,'CONSULTAR','Buscar una talla\nrequiere revisar\nel almacén.',CORAL)
card(4.77,2.45,3.8,'ORGANIZAR','Un encargo en papel\npuede perderse\nu olvidarse.',CORAL)
card(8.84,2.45,3.8,'REPONER','Sin datos reunidos,\ncuesta saber qué\nconviene comprar.',CORAL)
text(.85,5.75,11.8,.55,'No falta esfuerzo: falta información organizada.',24,INK,True)
finish('Estas son tres dificultades del caso. Hay que dedicar tiempo a buscar tallas, los encargos pueden quedar sin seguimiento y cuesta decidir qué reponer. La propuesta tiene sentido si resuelve esas tareas, no simplemente porque la tecnología parezca más moderna.')

new(5,'Después: una tienda conectada')
for i,(label,body) in enumerate([('WEB','Consultar y comprar'),('VENTAS','Registrar cada operación'),('INVENTARIO','Actualizar talla y color')]):
 x=.8+i*4.2; box(x,2.65,3.4,1.65,GREEN,r=.14); text(x+.3,2.98,2.9,.5,label,24,WHITE,True); text(x,4.7,3.8,.8,body,22)
 if i<2: arrow(x+3.55,3.2)
text(.8,5.95,11.8,.45,'La información debe circular entre los tres procesos.',23,GREEN,True)
finish('El cambio principal es conectar la página web, las ventas y el inventario. En una implantación real, registrar una venta debería actualizar las existencias disponibles. No se trata de llevar tres listas por separado, porque eso obligaría a copiar datos y podría crear nuevas diferencias.')

new(6,'La web: consultar desde casa')
browser(.8,2.5,6,3.55)
tag(7.5,2.55,'ANTES',CORAL); text(7.5,3.15,4.9,.85,'Ir o llamar para\nver qué hay.',23)
tag(7.5,4.35,'DESPUÉS'); text(7.5,4.95,4.95,.95,'Fotos, precios y tallas\nen un catálogo.',23)
finish('La web muestra el catálogo con fotografías o ilustraciones, precios, tallas y colores. El cliente puede consultar desde casa. En la demostración se puede buscar y filtrar, elegir una opción y añadirla a la cesta. La compra es simulada; no se realizan cobros ni envíos reales.')

new(7,'El stock: saber qué queda')
card(.8,2.45,5.3,'ANTES','Contar prendas a mano.\nRevisar el almacén\npara cada consulta.',CORAL)
card(7.15,2.45,5.3,'DESPUÉS','Consultar por talla y color.\nRegistrar entradas, ventas\ny devoluciones.',GREEN)
arrow(6.43,3.45)
text(.95,5.75,11.7,.55,'Mejora esperada: menos errores y consultas más rápidas.',22,GREEN,True)
finish('Cada prenda necesita un registro por talla y color. Por ejemplo, una camiseta negra M no puede confundirse con una L. Registrar entradas, ventas y devoluciones permite consultar las existencias. El sistema debe utilizarse bien: si se vende sin registrar la operación, la información dejará de ser fiable.')

new(8,'Pedidos y equipo: cada tarea, en su sitio')
for i,(label,body) in enumerate([('PEDIDO RECIBIDO','Queda registrado.'),('PREPARACIÓN','El personal sabe\nqué falta por hacer.'),('ENTREGA','Se actualiza\nel estado.')]):
 x=.8+i*4.2; text(x,2.5,3.8,.5,label,19,GREEN,True); box(x,3.25,3.4,1.3,WHITE,r=.15); text(x+.25,3.52,3,.9,body,21)
 if i<2: arrow(x+3.55,3.6)
box(.8,5.4,11.65,.8,PALE,r=.12); text(1.05,5.65,11.1,.35,'Un calendario compartido organiza turnos y responsabilidades.',20,INK,True)
finish('Los pedidos se guardarían con un estado para que el equipo sepa qué debe preparar o entregar. El calendario compartido serviría para turnos y tareas. Por ejemplo, se podría acordar quién recibe la mercancía y quién prepara los pedidos online. Estas funciones necesitan normas claras y formación.')

new(9,'Herramientas con una función concreta')
card(.7,2.5,3.8,'ODOO','Ventas, inventario\ny comercio\nelectrónico.')
card(4.77,2.5,3.8,'CANVA','Publicaciones\npara mostrar\nlas novedades.')
card(8.84,2.5,3.8,'CALENDARIO','Turnos, vacaciones\ny tareas\ncompartidas.')
text(.85,5.75,11.8,.65,'Primero se organizan los datos; después se forma al equipo.',22,INK,True)
finish('Odoo es una propuesta para integrar ventas, inventario y comercio electrónico. Canva puede servir para preparar contenido para redes sociales. Para los turnos bastaría inicialmente con un calendario compartido. La elección final dependería del presupuesto, de las funciones necesarias y de la facilidad de uso. Antes de empezar hay que ordenar los productos y formar al personal.')

new(10,'Qué enseña nuestra demostración')
card(.8,2.45,5.4,'SE PUEDE PROBAR','Catálogo y filtros.\nSelección de talla y color.\nCesta y pedido simulado.')
card(7.1,2.45,5.4,'PARA UNA TIENDA REAL','Cobros y envíos.\nDatos compartidos entre equipos.\nConexión con la gestión real.',CORAL)
text(.95,5.75,11.6,.7,'Los pedidos y el stock de la demo son ejemplos locales.',22,MUTED,True)
finish('La página es una demostración escolar. Permite enseñar el catálogo, las opciones de producto y una cesta con compra simulada. Los pedidos y el inventario de ejemplo se guardan en el navegador, si está disponible el almacenamiento local. No hemos conectado pagos, transportistas ni un sistema compartido entre varios equipos. Esas serían tareas de una implantación real.')

new(11,'Qué mejora y cómo lo comprobaríamos')
headers=[('COMODIDAD','Consultar desde casa','Revisar consultas y pedidos online.'),('TIEMPO','Encontrar una talla antes','Comparar el tiempo de consulta.'),('ORGANIZACIÓN','Menos olvidos y errores','Contar incidencias en pedidos.'),('REPOSICIÓN','Comprar con más información','Comparar stock registrado y real.')]
for i,(title,benefit,measure) in enumerate(headers):
 y=2.1+i*1.05; text(.85,y,2.55,.35,title,15,GREEN,True); text(3.25,y,4.25,.55,benefit,20,INK,True); text(8.0,y,4.6,.65,measure,17)
text(.85,6.48,11.8,.3,'Son mejoras esperadas. No se han medido resultados reales.',15,MUTED)
finish('Para comprobar la mejora compararíamos el tiempo de consulta de una talla, los errores de pedidos y las diferencias de inventario. También revisaríamos los pedidos procedentes de la web. Usaríamos periodos comparables y tendríamos en cuenta rebajas y temporadas. No ponemos porcentajes de mejora porque no se ha realizado una prueba real.')

new(12,'La transformación está en cómo trabajamos',dark=True)
text(.85,2.45,11.6,1.5,'Consultar mejor.\nOrganizar mejor. Atender mejor.',36,WHITE,True)
text(.85,4.35,11.6,.8,'La web abre un canal; la gestión digital conecta el negocio.',23,PALE)
text(.85,5.6,11.6,.25,'FUENTES DE LAS HERRAMIENTAS',11,PALE,True)
text(.85,5.98,11.6,.55,'odoo.com/es_ES/app/point-of-sale-features\ncanva.com/es_es/redes-sociales/',11,PALE)
finish('La conclusión es que digitalizar no consiste solo en abrir una web. La mejora aparece cuando la información de productos, ventas y pedidos está organizada y el equipo sabe utilizarla. La web ayuda al cliente; la gestión digital ayuda al personal. Hemos usado como referencia la documentación oficial de Odoo sobre punto de venta y la página oficial de Canva sobre contenido para redes sociales.')

cv.save(); prs.save(PPT)
with ZipFile(PPT) as z: assert z.testzip() is None
check=Presentation(PPT)
assert len(check.slides)==12
assert all(s.notes_slide.notes_text_frame.text.strip() for s in check.slides)
outside=[]
for i,s in enumerate(check.slides,1):
 for shp in s.shapes:
  if shp.left<0 or shp.top<0 or shp.left+shp.width>check.slide_width+100 or shp.top+shp.height>check.slide_height+100: outside.append((i,shp.name))
assert not outside, outside
doc=pdfium.PdfDocument(str(PDF)); assert len(doc)==12
thumbs=[]
for i in range(len(doc)):
 im=doc[i].render(scale=.65).to_pil().convert('RGB'); im.save(WORK/f'moda_slide_{i+1:02}.png'); thumbs.append(im)
sheet=Image.new('RGB',(3*624,4*351),(220,220,220))
for i,im in enumerate(thumbs): sheet.paste(im,(i%3*624,i//3*351))
sheet.save(WORK/'presentacion_contacto.jpg')
print(PPT); print(PDF); print('12 diapositivas; 12 notas de exposición; formas dentro del lienzo; ZIP y PDF correctos.')
