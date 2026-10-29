from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
from zipfile import ZipFile

out=Path(r'C:\Users\Javier\Documents\GitHub\Actividad-5\outputs')
out.mkdir(parents=True,exist_ok=True)
d=Document()
s=d.sections[0]
s.top_margin=Cm(1.8); s.bottom_margin=Cm(1.8); s.left_margin=Cm(2); s.right_margin=Cm(2)
s.page_width=Cm(21); s.page_height=Cm(29.7)
normal=d.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.line_spacing=1.12
for name in ['Title','Heading 1','Heading 2']:
 d.styles[name].font.name='Calibri'; d.styles[name].font.color.rgb=RGBColor.from_string('18545B')
d.styles['Heading 1'].font.size=Pt(17)
d.styles['Heading 2'].font.size=Pt(12)
h=s.header.paragraphs[0]; h.text='DIGITALIZACIÓN EN LOS SISTEMAS PRODUCTIVOS'; h.runs[0].font.size=Pt(9)
footer=s.footer.paragraphs[0]; footer.alignment=2
footer.add_run('Digitalización de una tienda de ropa · ')
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
def p(t): d.add_paragraph(t)
def head(t): d.add_heading(t,1)
def sub(t,txt): d.add_heading(t,2); p(txt)

d.add_heading('De la tienda física\na la tienda digital',0)
p('Propuesta de transformación digital de una tienda de ropa')
p('Caso ficticio: Moda Urbana')
head('1. El negocio y el objetivo')
p('Moda Urbana es una tienda ficticia de ropa y complementos. Para este trabajo, partimos de un pequeño comercio que vende en su local y organiza gran parte de sus tareas de forma manual. Esta situación nos permite explicar qué cambiaría al introducir herramientas digitales.')
p('El objetivo es mostrar cómo la transformación digital puede mejorar tanto el trabajo de la tienda como la experiencia del cliente. La propuesta principal es crear una página web, acompañada de un sistema para controlar las prendas y los pedidos. Los beneficios que se explican son esperados: no son resultados medidos en una empresa real.')
head('2. Antes: cómo funciona la tienda')
sub('Venta y atención al cliente','Los clientes tienen que acudir al local para conocer la colección, consultar precios y comprobar si queda su talla. También pueden llamar, pero el personal necesita revisar las prendas para responder. Las ventas dependen del horario de apertura y de las personas que visitan la zona.')
sub('Almacén y pedidos','Las entradas y salidas de ropa se apuntan en una libreta. Para saber cuántas unidades quedan de una talla o color, se revisa el almacén. Los encargos se anotan en papel y los pedidos a proveedores se realizan por teléfono. Un olvido puede provocar que las anotaciones no coincidan con las existencias reales.')
sub('Organización y publicidad','Los turnos y las tareas se comunican de palabra. La publicidad se basa en el escaparate y el boca a boca. Al no reunir los datos de ventas, resulta más difícil identificar qué prendas se venden mejor y decidir qué conviene reponer.')

d.add_page_break()
head('3. Después: qué proponemos mejorar')
sub('Una página web para consultar y comprar','La web tendría una página de inicio, un catálogo organizado por categorías y una ficha para cada prenda con fotografías, descripción, precio, tallas y colores. También incluiría información sobre la tienda, contacto, envíos y devoluciones. En una implantación real, el cliente podría completar una compra y recibir la confirmación de su pedido.')
sub('Un inventario conectado con las ventas','Cada combinación de prenda, talla y color tendría su propio registro. Al registrar una venta o recibir mercancía, se actualizarían las existencias. La tienda física y la web deberían compartir esa información para reducir el riesgo de ofrecer una prenda agotada. Las devoluciones también tendrían que registrarse.')
sub('Pedidos y reposición más organizados','Los pedidos de los clientes quedarían registrados con su estado: pendiente, preparado o enviado. El personal podría consultar qué falta por hacer. Además, revisar las existencias y las ventas ayudaría a preparar los pedidos a proveedores y a detectar las tallas que necesitan reposición.')
sub('Mejor coordinación del personal','Un calendario compartido permitiría consultar los turnos, las vacaciones y las tareas previstas. Por ejemplo, se podría asignar quién recibe la mercancía y quién prepara los pedidos online. El equipo necesitaría una formación inicial para utilizar las herramientas de forma correcta.')
sub('Más visibilidad y decisiones basadas en datos','Las redes sociales servirían para mostrar novedades y dirigir a los clientes hacia la web. Los registros de ventas permitirían comparar productos, tallas y temporadas. Así, las decisiones de compra se apoyarían en información del negocio y no solo en impresiones.')
head('4. Herramientas elegidas')
p('Odoo sería la herramienta principal para integrar el punto de venta, el inventario y el comercio electrónico. Su documentación incluye variantes de producto, control de existencias e integración entre ventas online y presenciales. Elegir un sistema conectado ayudaría a evitar que la misma información se copie en varios sitios.')
p('Canva se utilizaría para diseñar publicaciones de novedades y ofertas para las redes sociales. Para los turnos bastaría inicialmente con un calendario digital compartido. La elección definitiva dependería del presupuesto, las funciones necesarias y la facilidad de uso para el personal.')

d.add_page_break()
head('5. Comparación del antes y el después')
rows=[
('Ventas','Solo se compra en el local y durante su horario.','La web permite consultar y tramitar pedidos a cualquier hora.','Más comodidad y posibilidad de llegar a otros clientes.'),
('Consulta de prendas','Hay que acudir o llamar para preguntar.','El catálogo muestra fotos, precios y opciones de talla y color.','El cliente puede informarse desde casa.'),
('Inventario','Se cuentan prendas y se apuntan movimientos a mano.','Se registran las existencias por talla y color junto con las ventas.','Menos errores y consultas de stock más rápidas.'),
('Pedidos','Los encargos se anotan en papel.','Cada pedido tiene un registro y un estado.','Mejor seguimiento y menos olvidos.'),
('Personal','Los turnos y tareas se comunican de palabra.','Se consulta un calendario compartido.','Mayor claridad sobre horarios y responsabilidades.'),
('Publicidad','El escaparate y el boca a boca son los canales principales.','Las publicaciones en redes enlazan con la web.','Más visibilidad para las novedades.'),
('Reposición','Se decide principalmente por observación.','Se revisan los datos de ventas y las existencias.','Compras más ajustadas a la demanda.')]
t=d.add_table(rows=1, cols=4); t.style='Light Shading Accent 1'
for c,txt in zip(t.rows[0].cells,['Proceso','Antes','Después','Mejora esperada']): c.text=txt
for row in rows:
 for c,txt in zip(t.add_row().cells,row): c.text=txt
for row in t.rows:
 for cell in row.cells:
  for par in cell.paragraphs:
   par.paragraph_format.space_after=Pt(5)
   for r in par.runs: r.font.size=Pt(9)
 trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
head('6. Un ejemplo del cambio')
p('Antes, una persona pregunta por una camiseta negra de la talla M y el dependiente tiene que buscarla en el almacén. Si quiere encargarla, se anota su petición en papel. Después de la digitalización, podría consultar la ficha de la camiseta en la web y, si hay existencias, realizar el pedido. El personal lo vería registrado y podría prepararlo. La mejora está en conectar la consulta, la compra y el control de la prenda en un mismo proceso.')

d.add_page_break()
head('7. Cómo se pondría en marcha')
for title,txt in [
('Primero: organizar la información','Se haría un recuento de las prendas y se prepararía una ficha para cada producto, con sus tallas, colores, precios y fotografías.'),
('Después: preparar la web y la gestión','Se cargaría el catálogo y se configuraría la relación entre las ventas y el inventario. Se definirían los pasos para recibir, preparar y entregar los pedidos.'),
('Antes de utilizarlo: formar y comprobar','El personal practicaría cómo registrar ventas, entradas de mercancía y devoluciones. Se probaría un pedido completo para comprobar que la información se actualiza correctamente.'),
('Finalmente: revisar y ajustar','Se empezaría con un catálogo reducido y se ampliaría poco a poco. Se revisarían las incidencias para corregir los problemas antes de aumentar el volumen de pedidos.')]: sub(title,txt)
head('8. Cómo comprobaríamos la mejora')
p('Compararíamos el tiempo necesario para consultar una talla, las diferencias entre las existencias registradas y las reales, los errores en los pedidos y las ventas procedentes de la web. Utilizaríamos periodos de duración similar y tendríamos en cuenta si hay rebajas o cambios de temporada. Así podríamos valorar si las herramientas están resolviendo los problemas iniciales, sin dar por hecho que tener una web garantiza vender más.')
head('9. Conclusión')
p('La transformación digital de Moda Urbana no consiste únicamente en crear una página web. También supone cambiar cómo se organizan las ventas, el almacén y los pedidos. La web facilita el acceso a los productos, mientras que la gestión digital ayuda al personal a trabajar con información más ordenada. Si las herramientas se adaptan al negocio y el equipo aprende a utilizarlas, la tienda puede ofrecer un servicio más cómodo y reducir tareas manuales y errores.')
head('Fuentes de las herramientas')
p('Odoo. Funciones del punto de venta: inventario, variantes e integración con comercio electrónico. https://www.odoo.com/es_ES/app/point-of-sale-features')
p('Canva. Diseño de contenido para redes sociales. https://www.canva.com/es_es/redes-sociales/')
p('Este informe describe una propuesta. La web que se prepare para enseñarla en clase será una demostración; sus pantallas no acreditan por sí solas pagos, pedidos o inventario conectados a un sistema real.')
path=out/'Transformacion_digital_tienda_de_ropa.docx'
d.save(path)
with ZipFile(path) as z: assert z.testzip() is None
check=Document(path)
assert len(check.tables)==1 and len(check.tables[0].rows)==8
assert len(check.sections)==1
assert all(x in '\n'.join(p.text for p in check.paragraphs) for x in ['Conclusión','Caso ficticio','Fuentes de las herramientas'])
print(str(path))
print('Verificado: DOCX editable, 9 apartados, 7 procesos comparados y fuentes. Sin revisión visual de paginación.')
